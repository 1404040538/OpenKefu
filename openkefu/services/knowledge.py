from __future__ import annotations

import csv
import difflib
import hashlib
import io
import json
import logging
import re
import shutil
import threading
from collections import OrderedDict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import numpy as np

from openkefu.services import PROJECT_ROOT
from openkefu.services.embedding import build_embed_fn
from openkefu.web.config import AppConfig
from openkefu.web.db import Database, json_dumps

logger = logging.getLogger(__name__)


SUPPORTED_EXTENSIONS = {".txt", ".md", ".csv", ".xlsx", ".docx", ".pdf"}

# 店铺向量索引内存缓存：避免每条顾客消息都全量拉取 chunk 的
# embedding_json（LONGTEXT，每行可达数十 KB）并在 Python 里逐条解析。
INDEX_CACHE_MAX_SHOPS = 8


@dataclass
class _ShopIndex:
    fingerprint: tuple
    rows: list[dict[str, Any]] = field(default_factory=list)
    # (n, dim) float32，行向量已归一化，余弦相似度退化为点积。
    vectors: np.ndarray = field(default_factory=lambda: np.zeros((0, 0), dtype=np.float32))


class KnowledgeService:
    def __init__(self, db: Database, config: AppConfig):
        self.db = db
        self.config = config
        self._embed_fn = build_embed_fn(config.embedding)
        self._index_cache: OrderedDict[int, _ShopIndex] = OrderedDict()
        self._index_lock = threading.Lock()
        self._embed_config_warned = False

    def list_knowledge_bases(self) -> list[dict[str, Any]]:
        rows = self.db.query(
            """
            SELECT kb.*,
                   COUNT(DISTINCT kbs.shop_id) AS shop_count,
                   COUNT(DISTINCT kf.id) AS file_count
            FROM knowledge_bases kb
            LEFT JOIN knowledge_base_shops kbs ON kbs.knowledge_base_id=kb.id
            LEFT JOIN knowledge_files kf ON kf.knowledge_base_id=kb.id
            GROUP BY kb.id
            ORDER BY kb.id DESC
            """
        )
        for row in rows:
            row["shop_ids"] = [
                int(item["shop_id"])
                for item in self.db.query(
                    "SELECT shop_id FROM knowledge_base_shops WHERE knowledge_base_id=%s ORDER BY shop_id",
                    (row["id"],),
                )
            ]
        return rows

    def create_knowledge_base(self, *, name: str, description: str, shop_ids: list[int], created_by: int | None) -> dict[str, Any]:
        kb_id = self.db.execute(
            """
            INSERT INTO knowledge_bases (name, description, created_by)
            VALUES (%s,%s,%s)
            """,
            (name.strip(), description.strip(), created_by),
        )
        self.set_shop_bindings(kb_id, shop_ids)
        return self.get_knowledge_base(kb_id)

    def get_knowledge_base(self, kb_id: int) -> dict[str, Any]:
        row = self.db.query_one("SELECT * FROM knowledge_bases WHERE id=%s", (kb_id,))
        if not row:
            raise ValueError(f"knowledge base not found: {kb_id}")
        row["shop_ids"] = [
            int(item["shop_id"])
            for item in self.db.query(
                "SELECT shop_id FROM knowledge_base_shops WHERE knowledge_base_id=%s ORDER BY shop_id",
                (kb_id,),
            )
        ]
        row["files"] = self.db.query(
            "SELECT * FROM knowledge_files WHERE knowledge_base_id=%s ORDER BY id DESC",
            (kb_id,),
        )
        row["qa_items"] = self.list_qa_items(kb_id)
        return row

    def update_knowledge_base(
        self,
        kb_id: int,
        *,
        name: str | None = None,
        description: str | None = None,
        status: str | None = None,
        shop_ids: list[int] | None = None,
    ) -> dict[str, Any]:
        fields = []
        values: list[Any] = []
        if name is not None:
            fields.append("name=%s")
            values.append(name.strip())
        if description is not None:
            fields.append("description=%s")
            values.append(description.strip())
        if status is not None:
            fields.append("status=%s")
            values.append(status)
        if fields:
            values.append(kb_id)
            self.db.execute(f"UPDATE knowledge_bases SET {', '.join(fields)} WHERE id=%s", values)
        if shop_ids is not None:
            self.set_shop_bindings(kb_id, shop_ids)
        return self.get_knowledge_base(kb_id)

    def delete_knowledge_base(self, kb_id: int) -> None:
        row = self.db.query_one("SELECT * FROM knowledge_bases WHERE id=%s", (kb_id,))
        if not row:
            raise ValueError(f"knowledge base not found: {kb_id}")
        files = self.db.query("SELECT file_path FROM knowledge_files WHERE knowledge_base_id=%s", (kb_id,))
        self.db.execute("DELETE FROM knowledge_bases WHERE id=%s", (kb_id,))
        for item in files:
            path = Path(str(item.get("file_path") or ""))
            try:
                if path.exists() and path.is_file():
                    path.unlink()
            except OSError:
                pass
        kb_dir = self._knowledge_dir() / f"kb_{kb_id}"
        try:
            if kb_dir.exists():
                shutil.rmtree(kb_dir)
        except OSError:
            pass

    def set_shop_bindings(self, kb_id: int, shop_ids: list[int]) -> None:
        self.db.execute("DELETE FROM knowledge_base_shops WHERE knowledge_base_id=%s", (kb_id,))
        for shop_id in sorted({int(item) for item in shop_ids}):
            self.db.execute(
                "INSERT IGNORE INTO knowledge_base_shops (knowledge_base_id, shop_id) VALUES (%s,%s)",
                (kb_id, shop_id),
            )

    def list_qa_items(self, kb_id: int) -> list[dict[str, Any]]:
        return self.db.query(
            """
            SELECT id, knowledge_base_id, question, reply, image_base64, image_name,
                   image_content_type, sort_order, created_at, updated_at
            FROM knowledge_qa_items
            WHERE knowledge_base_id=%s
            ORDER BY sort_order ASC, id ASC
            """,
            (kb_id,),
        )

    def set_qa_items(self, kb_id: int, items: list[dict[str, Any]], created_by: int | None = None) -> list[dict[str, Any]]:
        if not self.db.query_one("SELECT id FROM knowledge_bases WHERE id=%s", (kb_id,)):
            raise ValueError(f"knowledge base not found: {kb_id}")

        cleaned: list[dict[str, Any]] = []
        for index, item in enumerate(items):
            question = str(item.get("question") or "").strip()
            reply = str(item.get("reply") or "").strip()
            image_base64 = str(item.get("image_base64") or "").strip()
            image_name = str(item.get("image_name") or "").strip()
            image_content_type = str(item.get("image_content_type") or "").strip()
            if not question:
                raise ValueError(f"第 {index + 1} 行问题不能为空")
            if not reply and not image_base64:
                raise ValueError(f"第 {index + 1} 行回复和图片至少填写一项")
            if image_base64:
                if not image_base64.startswith("data:image/"):
                    raise ValueError(f"第 {index + 1} 行图片格式不正确")
                image_payload = image_base64.split(",", 1)[1] if "," in image_base64 else image_base64
                if len(image_payload.encode("utf-8")) > 5 * 1024 * 1024:
                    raise ValueError(f"第 {index + 1} 行图片大小不能超过 5MB")
            cleaned.append(
                {
                    "question": question,
                    "reply": reply,
                    "image_base64": image_base64,
                    "image_name": image_name,
                    "image_content_type": image_content_type,
                    "sort_order": index,
                }
            )

        with self.db.connect() as conn:
            with conn.cursor() as cursor:
                cursor.execute("DELETE FROM knowledge_qa_items WHERE knowledge_base_id=%s", (kb_id,))
                for item in cleaned:
                    cursor.execute(
                        """
                        INSERT INTO knowledge_qa_items
                        (knowledge_base_id, question, reply, image_base64, image_name, image_content_type, sort_order)
                        VALUES (%s,%s,%s,%s,%s,%s,%s)
                        """,
                        (
                            kb_id,
                            item["question"],
                            item["reply"] or None,
                            item["image_base64"] or None,
                            item["image_name"] or None,
                            item["image_content_type"] or None,
                            item["sort_order"],
                        ),
                    )
                cursor.execute("UPDATE knowledge_bases SET updated_at=NOW() WHERE id=%s", (kb_id,))
        return self.list_qa_items(kb_id)

    def upload_file_bytes(self, *, kb_id: int, filename: str, content_type: str | None, data: bytes) -> dict[str, Any]:
        suffix = Path(filename).suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise ValueError(f"unsupported file type: {suffix}")
        storage_dir = self._knowledge_dir() / f"kb_{kb_id}"
        storage_dir.mkdir(parents=True, exist_ok=True)
        safe_name = _safe_filename(filename)
        digest = hashlib.sha256(data).hexdigest()
        file_id = self.db.execute(
            """
            INSERT INTO knowledge_files
            (knowledge_base_id, filename, content_type, file_path, file_size, file_hash, status)
            VALUES (%s,%s,%s,%s,%s,%s,'uploaded')
            """,
            (kb_id, filename, content_type or "", "", len(data), digest),
        )
        file_path = storage_dir / f"{file_id}_{safe_name}"
        file_path.write_bytes(data)
        self.db.execute("UPDATE knowledge_files SET file_path=%s WHERE id=%s", (str(file_path), file_id))
        self.ingest_file(file_id)
        return self.db.query_one("SELECT * FROM knowledge_files WHERE id=%s", (file_id,))

    def ingest_file(self, file_id: int) -> None:
        file_row = self.db.query_one("SELECT * FROM knowledge_files WHERE id=%s", (file_id,))
        if not file_row:
            raise ValueError(f"knowledge file not found: {file_id}")
        try:
            self.db.execute("UPDATE knowledge_files SET status='processing', error=NULL WHERE id=%s", (file_id,))
            sections = self._extract_sections(Path(file_row["file_path"]))
            chunks = []
            for section_text, source_label in sections:
                for text in _chunk_text(section_text):
                    chunks.append((text, source_label))
            if not chunks:
                raise ValueError("文件没有解析出可用文本")

            kb_id = file_row["knowledge_base_id"]
            # vector_id 由 file_id+chunk_index 确定，无需插入后再逐行回填。
            insert_values = [
                (kb_id, file_id, index, content, source_label, f"chunk:{file_id}:{index}")
                for index, (content, source_label) in enumerate(chunks)
            ]
            self.db.execute("DELETE FROM knowledge_chunks WHERE file_id=%s", (file_id,))
            self.db.execute_many(
                """
                INSERT INTO knowledge_chunks
                (knowledge_base_id, file_id, chunk_index, content, source_label, vector_id)
                VALUES (%s,%s,%s,%s,%s,%s)
                """,
                insert_values,
            )

            embeddings = self._embed([content for content, _ in chunks])
            if len(embeddings) != len(chunks):
                raise RuntimeError("embedding result count does not match knowledge chunks")
            current_model = self.config.embedding.model
            self.db.execute_many(
                "UPDATE knowledge_chunks SET embedding_json=%s, embedding_model=%s WHERE file_id=%s AND chunk_index=%s",
                [
                    (json_dumps(embedding), current_model, file_id, index)
                    for index, embedding in enumerate(embeddings)
                ],
            )
            self.db.execute("UPDATE knowledge_files SET status='ready', error=NULL WHERE id=%s", (file_id,))
        except Exception as exc:
            self.db.execute("UPDATE knowledge_files SET status='failed', error=%s WHERE id=%s", (str(exc), file_id))
            raise

    def search_for_shop(self, *, shop_id: int, query: str, top_k: int = 5) -> list[dict[str, Any]]:
        query = query.strip()
        if not query:
            return []
        if self._embed_fn is None:
            if not self._embed_config_warned:
                self._embed_config_warned = True
                logger.warning("embedding 未配置，知识库语义检索不可用（QA 问答与注意事项不受影响）")
            return []
        try:
            index = self._shop_index(shop_id)
            if index is None or not index.rows or index.vectors.size == 0:
                return []
            query_embedding = np.asarray(self._embed_fn([query[:1000]])[0], dtype=np.float32)
            query_norm = float(np.linalg.norm(query_embedding))
            if not query_norm:
                return []
            query_embedding = query_embedding / query_norm
            similarities = index.vectors @ query_embedding
            order = np.argsort(-similarities)[:top_k]
            hits: list[dict[str, Any]] = []
            for position in order:
                similarity = float(similarities[position])
                row = dict(index.rows[int(position)])
                row["distance"] = 1.0 - max(-1.0, min(1.0, similarity))
                hits.append(row)
            return hits
        except Exception:
            logger.debug("knowledge search failed", exc_info=True)
            return []

    def _shop_index(self, shop_id: int) -> _ShopIndex | None:
        fingerprint = self._index_fingerprint(shop_id)
        with self._index_lock:
            cached = self._index_cache.get(shop_id)
            if cached is not None and cached.fingerprint == fingerprint:
                self._index_cache.move_to_end(shop_id)
                return cached

        rows = self._eligible_chunks(shop_id)
        if self._ensure_chunk_embeddings(rows):
            # 回填会改变 missing/stale 计数，指纹需要重算。
            fingerprint = self._index_fingerprint(shop_id)

        vectors: list[np.ndarray] = []
        meta_rows: list[dict[str, Any]] = []
        dimension = 0
        for row in rows:
            try:
                embedding = np.asarray(json.loads(row.get("embedding_json") or "[]"), dtype=np.float32)
            except (TypeError, ValueError, json.JSONDecodeError):
                continue
            if embedding.ndim != 1 or embedding.size == 0:
                continue
            if dimension == 0:
                dimension = int(embedding.size)
            elif embedding.size != dimension:
                continue
            norm = float(np.linalg.norm(embedding))
            if not norm:
                continue
            meta = dict(row)
            meta.pop("embedding_json", None)
            meta_rows.append(meta)
            vectors.append(embedding / norm)

        index = _ShopIndex(
            fingerprint=fingerprint,
            rows=meta_rows,
            vectors=(np.vstack(vectors).astype(np.float32) if vectors else np.zeros((0, 0), dtype=np.float32)),
        )
        with self._index_lock:
            self._index_cache[shop_id] = index
            self._index_cache.move_to_end(shop_id)
            while len(self._index_cache) > INDEX_CACHE_MAX_SHOPS:
                self._index_cache.popitem(last=False)
        return index

    def _index_fingerprint(self, shop_id: int) -> tuple:
        row = self.db.query_one(
            """
            SELECT COUNT(*) AS cnt,
                   COALESCE(MAX(kc.id), 0) AS max_id,
                   COALESCE(SUM(kc.embedding_json IS NULL), 0) AS missing_embedding,
                   COALESCE(SUM(kc.embedding_model IS NULL OR kc.embedding_model <> %s), 0) AS stale_model
            FROM knowledge_chunks kc
            JOIN knowledge_files kf ON kf.id=kc.file_id
            JOIN knowledge_bases kb ON kb.id=kc.knowledge_base_id
            JOIN knowledge_base_shops kbs ON kbs.knowledge_base_id=kb.id
            WHERE kbs.shop_id=%s AND kb.status='active' AND kf.status='ready'
            """,
            (self.config.embedding.model, shop_id),
        )
        if not row:
            return (0, 0, 0, 0)
        return (
            int(row.get("cnt") or 0),
            int(row.get("max_id") or 0),
            int(row.get("missing_embedding") or 0),
            int(row.get("stale_model") or 0),
        )

    def find_qa_match_for_shop(self, *, shop_id: int, query: str) -> dict[str, Any] | None:
        normalized_query = _normalize_question(query)
        if not normalized_query:
            return None
        rows = self.db.query(
            """
            SELECT kqa.*, kb.name AS knowledge_base_name
            FROM knowledge_qa_items kqa
            JOIN knowledge_bases kb ON kb.id=kqa.knowledge_base_id
            JOIN knowledge_base_shops kbs ON kbs.knowledge_base_id=kb.id
            WHERE kbs.shop_id=%s AND kb.status='active'
            ORDER BY kqa.sort_order ASC, kqa.id ASC
            """,
            (shop_id,),
        )
        best: tuple[float, dict[str, Any]] | None = None
        for row in rows:
            question = _normalize_question(str(row.get("question") or ""))
            if not question:
                continue
            if normalized_query == question:
                score = 1.0
            elif question in normalized_query or normalized_query in question:
                score = 0.92
            else:
                score = difflib.SequenceMatcher(None, normalized_query, question).ratio()
            if best is None or score > best[0]:
                best = (score, row)
        if not best or best[0] < 0.58:
            return None
        matched = dict(best[1])
        matched["match_score"] = best[0]
        return matched

    def _eligible_chunks(self, shop_id: int) -> list[dict[str, Any]]:
        return self.db.query(
            """
            SELECT kc.*, kf.filename, kb.name AS knowledge_base_name
            FROM knowledge_chunks kc
            JOIN knowledge_files kf ON kf.id=kc.file_id
            JOIN knowledge_bases kb ON kb.id=kc.knowledge_base_id
            JOIN knowledge_base_shops kbs ON kbs.knowledge_base_id=kb.id
            WHERE kbs.shop_id=%s
              AND kb.status='active'
              AND kf.status='ready'
            """,
            (shop_id,),
        )

    def _ensure_chunk_embeddings(self, chunks: list[dict[str, Any]], batch_size: int = 64) -> bool:
        """回填缺失或模型已更换的向量；返回是否有回填。

        按 embedding_model 判断失效：换 embedding 模型后旧向量维度不同，
        直接跳过会导致检索永远为空，必须整体重算。
        """
        current_model = self.config.embedding.model
        missing = [
            item
            for item in chunks
            if not item.get("embedding_json") or str(item.get("embedding_model") or "") != current_model
        ]
        for start in range(0, len(missing), batch_size):
            batch = missing[start:start + batch_size]
            embeddings = self._embed([str(item.get("content") or "") for item in batch])
            if len(embeddings) != len(batch):
                raise RuntimeError("embedding backfill result count does not match knowledge chunks")
            for item, embedding in zip(batch, embeddings):
                encoded = json_dumps(embedding)
                self.db.execute(
                    "UPDATE knowledge_chunks SET embedding_json=%s, embedding_model=%s WHERE id=%s",
                    (encoded, current_model, item["id"]),
                )
                item["embedding_json"] = encoded
                item["embedding_model"] = current_model
        return bool(missing)

    def _embed(self, texts: list[str]) -> list[list[float]]:
        if self._embed_fn is None:
            raise RuntimeError("缺少 embedding 配置，请配置 embedding.api_key/base_url/model")
        return self._embed_fn(texts)

    def _knowledge_dir(self) -> Path:
        path = Path(self.config.storage.knowledge_dir)
        return path if path.is_absolute() else PROJECT_ROOT / path

    def _extract_sections(self, path: Path) -> list[tuple[str, str]]:
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md"}:
            return [(path.read_text(encoding="utf-8", errors="ignore"), path.name)]
        if suffix == ".csv":
            return [(_read_csv(path), path.name)]
        if suffix == ".xlsx":
            return _read_xlsx(path)
        if suffix == ".docx":
            return [(_read_docx(path), path.name)]
        if suffix == ".pdf":
            return _read_pdf(path)
        raise ValueError(f"unsupported file type: {suffix}")


class NoteSetService:
    def __init__(self, db: Database):
        self.db = db

    def list_note_sets(self) -> list[dict[str, Any]]:
        rows = self.db.query(
            """
            SELECT ns.*,
                   COUNT(DISTINCT nss.shop_id) AS shop_count,
                   COUNT(DISTINCT nsi.id) AS item_count
            FROM note_sets ns
            LEFT JOIN note_set_shops nss ON nss.note_set_id=ns.id
            LEFT JOIN note_set_items nsi ON nsi.note_set_id=ns.id
            GROUP BY ns.id
            ORDER BY ns.id DESC
            """
        )
        for row in rows:
            row["shop_ids"] = [
                int(item["shop_id"])
                for item in self.db.query(
                    "SELECT shop_id FROM note_set_shops WHERE note_set_id=%s ORDER BY shop_id",
                    (row["id"],),
                )
            ]
        return rows

    def create_note_set(self, *, name: str, description: str, shop_ids: list[int], created_by: int | None) -> dict[str, Any]:
        ns_id = self.db.execute(
            "INSERT INTO note_sets (name, description, created_by) VALUES (%s,%s,%s)",
            (name.strip(), description.strip(), created_by),
        )
        self.set_shop_bindings(ns_id, shop_ids)
        return self.get_note_set(ns_id)

    def get_note_set(self, ns_id: int) -> dict[str, Any]:
        row = self.db.query_one("SELECT * FROM note_sets WHERE id=%s", (ns_id,))
        if not row:
            raise ValueError(f"note set not found: {ns_id}")
        row["shop_ids"] = [
            int(item["shop_id"])
            for item in self.db.query(
                "SELECT shop_id FROM note_set_shops WHERE note_set_id=%s ORDER BY shop_id",
                (ns_id,),
            )
        ]
        row["items"] = self.db.query(
            "SELECT * FROM note_set_items WHERE note_set_id=%s ORDER BY id ASC",
            (ns_id,),
        )
        return row

    def update_note_set(self, ns_id: int, *, name: str | None = None, description: str | None = None, status: str | None = None, shop_ids: list[int] | None = None) -> dict[str, Any]:
        fields = []
        values: list[Any] = []
        if name is not None:
            fields.append("name=%s")
            values.append(name.strip())
        if description is not None:
            fields.append("description=%s")
            values.append(description.strip())
        if status is not None:
            fields.append("status=%s")
            values.append(status)
        if fields:
            values.append(ns_id)
            self.db.execute(f"UPDATE note_sets SET {', '.join(fields)} WHERE id=%s", values)
        if shop_ids is not None:
            self.set_shop_bindings(ns_id, shop_ids)
        return self.get_note_set(ns_id)

    def delete_note_set(self, ns_id: int) -> None:
        self.db.execute("DELETE FROM note_sets WHERE id=%s", (ns_id,))

    def set_shop_bindings(self, ns_id: int, shop_ids: list[int]) -> None:
        self.db.execute("DELETE FROM note_set_shops WHERE note_set_id=%s", (ns_id,))
        for shop_id in sorted({int(item) for item in shop_ids}):
            self.db.execute(
                "INSERT IGNORE INTO note_set_shops (note_set_id, shop_id) VALUES (%s,%s)",
                (ns_id, shop_id),
            )

    def add_item(self, ns_id: int, content: str, created_by: int | None = None) -> dict[str, Any]:
        item_id = self.db.execute(
            "INSERT INTO note_set_items (note_set_id, content, created_by) VALUES (%s,%s,%s)",
            (ns_id, content.strip(), created_by),
        )
        self.db.execute("UPDATE note_sets SET updated_at=NOW() WHERE id=%s", (ns_id,))
        return self.db.query_one("SELECT * FROM note_set_items WHERE id=%s", (item_id,))

    def update_item(self, item_id: int, content: str) -> dict[str, Any]:
        self.db.execute("UPDATE note_set_items SET content=%s WHERE id=%s", (content.strip(), item_id))
        row = self.db.query_one("SELECT note_set_id FROM note_set_items WHERE id=%s", (item_id,))
        if row:
            self.db.execute("UPDATE note_sets SET updated_at=NOW() WHERE id=%s", (row["note_set_id"],))
        return self.db.query_one("SELECT * FROM note_set_items WHERE id=%s", (item_id,))

    def delete_item(self, item_id: int) -> None:
        row = self.db.query_one("SELECT note_set_id FROM note_set_items WHERE id=%s", (item_id,))
        self.db.execute("DELETE FROM note_set_items WHERE id=%s", (item_id,))
        if row:
            self.db.execute("UPDATE note_sets SET updated_at=NOW() WHERE id=%s", (row["note_set_id"],))

    def get_items_for_shop(self, shop_id: int) -> list[str]:
        rows = self.db.query(
            """
            SELECT nsi.content
            FROM note_set_items nsi
            JOIN note_sets ns ON ns.id=nsi.note_set_id
            JOIN note_set_shops nss ON nss.note_set_id=ns.id
            WHERE nss.shop_id=%s AND ns.status='active'
            ORDER BY ns.id, nsi.id ASC
            """,
            (shop_id,),
        )
        return [row["content"] for row in rows]


def _safe_filename(filename: str) -> str:
    name = Path(filename).name
    return re.sub(r"[^A-Za-z0-9._\-\u4e00-\u9fff]+", "_", name)[:180] or "upload"


def _normalize_question(text: str) -> str:
    return re.sub(r"[\s\?？!！,，.。:：;；]+", "", text).lower()


def _chunk_text(text: str, *, size: int = 1000, overlap: int = 120) -> list[str]:
    cleaned = re.sub(r"\s+", " ", text).strip()
    if not cleaned:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(cleaned):
        end = min(len(cleaned), start + size)
        chunks.append(cleaned[start:end])
        if end >= len(cleaned):
            break
        start = max(0, end - overlap)
    return chunks


def _read_csv(path: Path) -> str:
    data = path.read_bytes()
    text = data.decode("utf-8-sig", errors="ignore")
    rows = csv.reader(io.StringIO(text))
    return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)


def _read_xlsx(path: Path) -> list[tuple[str, str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sections: list[tuple[str, str]] = []
    for sheet in workbook.worksheets:
        lines = []
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(value.strip() for value in values):
                lines.append(" | ".join(values))
        if lines:
            sections.append(("\n".join(lines), f"{path.name}:{sheet.title}"))
    return sections


def _read_docx(path: Path) -> str:
    from docx import Document

    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _read_pdf(path: Path) -> list[tuple[str, str]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    sections: list[tuple[str, str]] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            sections.append((text, f"{path.name}:page-{index}"))
    return sections
